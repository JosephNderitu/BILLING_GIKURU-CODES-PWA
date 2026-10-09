import calendar
from datetime import date

from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Client, Document
from decimal import Decimal

ZERO = Decimal("0")


def sales_q():
    """Same rule as Document.is_sale, expressed for the database."""
    return (Q(doc_type="invoice", status__in=["sent", "paid"]) |
            (Q(doc_type="receipt", status="paid") & ~Q(source__doc_type="invoice")))


def sales_docs(start, end):
    return (Document.objects.filter(sales_q(), issue_date__gte=start, issue_date__lte=end)
            .select_related("client").prefetch_related("items").order_by("issue_date", "id"))


def _date(value, default):
    try:
        return parse_date(value or "") or default
    except ValueError:
        return default


def _qs(request):
    params = request.GET.copy()
    params.pop("export", None)
    return params.urlencode()


def workbook_response(filename, sheets):
    """sheets: [(title, headers, rows, money_column_numbers)]"""
    wb = Workbook()
    wb.remove(wb.active)
    fill = PatternFill("solid", fgColor="15803D")
    for title, headers, rows, money_cols in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), fill, Alignment(vertical="center")
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, 1):
            longest = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), 45)
        for col in money_cols:
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                for cell in row:
                    cell.number_format = "#,##0.00"
    resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(resp)
    return resp


# ------------------------------------------------------------------ sales by month
@login_required
def sales(request):
    today = timezone.localdate()
    try:
        year = int(request.GET.get("year", today.year))
    except ValueError:
        year = today.year
    months = {m: {"count": 0, "net": ZERO, "vat": ZERO, "total": ZERO} for m in range(1, 13)}
    for d in sales_docs(date(year, 1, 1), date(year, 12, 31)):
        r = months[d.issue_date.month]
        r["count"] += 1
        r["net"] += d.net_amount
        r["vat"] += d.vat_amount
        r["total"] += d.total
    peak = max((r["total"] for r in months.values()), default=ZERO) or Decimal("1")
    rows = [{"label": calendar.month_name[m], **months[m], "pct": int(months[m]["total"] / peak * 100)}
            for m in range(1, 13)]
    totals = {k: sum((r[k] for r in rows), ZERO) for k in ("count", "net", "vat", "total")}

    if request.GET.get("export") == "xlsx":
        data = [[r["label"], r["count"], r["net"], r["vat"], r["total"]] for r in rows]
        data.append(["Total", totals["count"], totals["net"], totals["vat"], totals["total"]])
        return workbook_response(f"sales-{year}.xlsx", [
            (f"Sales {year}", ["Month", "Documents", "Net", "VAT", "Total"], data, [3, 4, 5])])

    years = sorted({d.year for d in Document.objects.dates("issue_date", "year")} | {today.year}, reverse=True)
    return render(request, "documents/reports/sales.html", {
        "tab": "sales", "rows": rows, "totals": totals, "year": year, "years": years,
        "avg": (totals["total"] / totals["count"]) if totals["count"] else ZERO})


# ------------------------------------------------------------------ VAT summary
@login_required
def vat(request):
    today = timezone.localdate()
    start = _date(request.GET.get("from"), today.replace(day=1))
    end = _date(request.GET.get("to"), today)
    docs = list(sales_docs(start, end))
    taxable = vat_total = exempt = gross = ZERO
    by_month = {}
    for d in docs:
        m = by_month.setdefault(d.issue_date.strftime("%Y-%m"), {"net": ZERO, "vat": ZERO, "none": ZERO})
        if d.vat_mode == "none":
            exempt += d.total
            m["none"] += d.total
        else:
            taxable += d.net_amount
            vat_total += d.vat_amount
            m["net"] += d.net_amount
            m["vat"] += d.vat_amount
        gross += d.total

    if request.GET.get("export") == "xlsx":
        summary = [["Period", f"{start} to {end}"], ["Taxable sales (net)", taxable],
                   ["Output VAT", vat_total], ["Sales without VAT", exempt], ["Total sales", gross]]
        detail = [[d.issue_date, d.number, d.get_doc_type_display(), d.display_client,
                   d.buyer_pin_display, d.etims_invoice_no, d.get_vat_mode_display(),
                   d.net_amount, d.vat_amount, d.total] for d in docs]
        return workbook_response(f"vat-{start}-{end}.xlsx", [
            ("Summary", ["Item", "Amount"], summary, [2]),
            ("Documents", ["Date", "Number", "Type", "Client", "Buyer PIN", "eTIMS no.", "VAT mode",
                           "Net", "VAT", "Total"], detail, [8, 9, 10])])

    return render(request, "documents/reports/vat.html", {
        "tab": "vat", "docs": docs, "start": start, "end": end, "taxable": taxable,
        "vat_total": vat_total, "exempt": exempt, "gross": gross,
        "months": sorted(by_month.items()), "querystring": _qs(request)})


# ------------------------------------------------------------------ aging
LABELS = ["Not yet due", "1-30 days", "31-60 days", "61-90 days", "Over 90 days"]


def _bucket(days):
    return 0 if days <= 0 else 1 if days <= 30 else 2 if days <= 60 else 3 if days <= 90 else 4


@login_required
def aging(request):
    today = timezone.localdate()
    invoices = (Document.objects.filter(doc_type="invoice", status="sent")
                .select_related("client").prefetch_related("items"))
    by_client, detail, totals = {}, [], [ZERO] * 5
    for d in invoices:
        due = d.due_date or d.issue_date
        days = (today - due).days
        i, amt = _bucket(days), d.total
        by_client.setdefault(d.display_client, [ZERO] * 5)[i] += amt
        totals[i] += amt
        detail.append({"doc": d, "due": due, "days": max(days, 0), "bucket": LABELS[i], "amount": amt})
    detail.sort(key=lambda r: -r["days"])
    clients = sorted(({"name": n, "cells": c, "total": sum(c, ZERO)} for n, c in by_client.items()),
                     key=lambda r: -r["total"])

    if request.GET.get("export") == "xlsx":
        return workbook_response(f"aging-{today}.xlsx", [
            ("By client", ["Client"] + LABELS + ["Total"],
             [[c["name"]] + c["cells"] + [c["total"]] for c in clients]
             + [["Total"] + totals + [sum(totals, ZERO)]], [2, 3, 4, 5, 6, 7]),
            ("Invoices", ["Number", "Client", "Issued", "Due", "Days overdue", "Bucket", "Amount"],
             [[r["doc"].number, r["doc"].display_client, r["doc"].issue_date, r["due"],
               r["days"], r["bucket"], r["amount"]] for r in detail], [7])])

    return render(request, "documents/reports/aging.html", {
        "tab": "aging", "labels": LABELS, "buckets": list(zip(LABELS, totals)), "clients": clients,
        "totals": totals, "grand": sum(totals, ZERO), "detail": detail, "today": today})


# ------------------------------------------------------------------ client statements
def build_statement(client, start, end):
    docs = list(Document.objects.filter(client=client).exclude(status="void")
                .select_related("source").prefetch_related("items"))
    receipted = {d.source_id for d in docs if d.doc_type == "receipt" and d.source_id}
    ledger = []  # (date, created, text, debit, credit, doc)
    for d in docs:
        if d.doc_type == "invoice" and d.status in ("sent", "paid"):
            ledger.append((d.issue_date, d.created_at, f"Invoice {d.number}", d.total, ZERO, d))
            if d.status == "paid" and d.pk not in receipted:
                ledger.append((d.issue_date, d.created_at, f"Payment for {d.number} (marked paid)", ZERO, d.total, d))
        elif d.doc_type == "receipt" and d.status == "paid":
            if d.source_id:
                ledger.append((d.issue_date, d.created_at, f"Receipt {d.number}", ZERO, d.total, d))
            else:
                ledger.append((d.issue_date, d.created_at, f"Cash sale {d.number}", d.total, d.total, d))
    ledger.sort(key=lambda r: (r[0], r[1]))

    balance, opening, rows = ZERO, None, []
    for when, _, text, debit, credit, doc in ledger:
        if when >= start and opening is None:
            opening = balance
        balance += debit - credit
        if start <= when <= end:
            rows.append({"date": when, "text": text, "debit": debit, "credit": credit,
                         "balance": balance, "doc": doc})
    if opening is None:
        opening = balance
    return {"rows": rows, "opening": opening, "closing": rows[-1]["balance"] if rows else opening,
            "debit": sum((r["debit"] for r in rows), ZERO), "credit": sum((r["credit"] for r in rows), ZERO)}


@login_required
def statements(request):
    today = timezone.localdate()
    clients = Client.objects.all()
    cid = request.GET.get("client", "")
    client = clients.filter(pk=cid).first() if cid.isdigit() else None
    start = _date(request.GET.get("from"), date(today.year, 1, 1))
    end = _date(request.GET.get("to"), today)
    data = build_statement(client, start, end) if client else None

    if data and request.GET.get("export") == "xlsx":
        rows = [["", "Opening balance", "", "", data["opening"]]]
        rows += [[r["date"], r["text"], r["debit"], r["credit"], r["balance"]] for r in data["rows"]]
        rows.append(["", "Closing balance", data["debit"], data["credit"], data["closing"]])
        return workbook_response(f"statement-{client.name}-{end}.xlsx".replace(" ", "_"), [
            ("Statement", ["Date", "Description", "Debit", "Credit", "Balance"], rows, [3, 4, 5])])

    return render(request, "documents/reports/statements.html", {
        "tab": "statements", "clients": clients, "client": client, "data": data,
        "start": start, "end": end, "querystring": _qs(request)})